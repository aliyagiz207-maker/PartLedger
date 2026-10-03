import pandas as pd
from openpyxl import load_workbook
from src.report_generator import generate_report

def test_generate_report_creates_excel(tmp_path):
    df = pd.DataFrame({
        "Date": pd.to_datetime(["2026-01-01"]),
        "Product": ["Brake Pad"],
        "Region": ["İzmir"],
        "Quantity": [10],
        "Revenue": [1000],
        "Profit": [400],
        "Margin": [40.0],
    })
    kpis = {"Total Quantity": 10, "Total Revenue": 1000, "Profit Margin": 40.0}
    region_summary = pd.DataFrame({"Region": ["İzmir"], "Revenue": [1000], "Profit": [400], "Margin": [40.0]})
    product_summary = pd.DataFrame({"Product": ["Brake Pad"], "Revenue": [1000], "Profit": [400], "Margin": [40.0]})
    monthly_summary = pd.DataFrame({"Month": ["2026-01"], "Revenue": [1000], "Profit": [400]})
    
    test_output = tmp_path / "Test_Dashboard.xlsx"
    config = {
        "output_file": str(test_output),
        "logo_path": "non_existent_logo.png",
        "dashboard_title": "PartLedger Test Dashboard",
        "company_name": "Terakki Test",
        "currency": "₺"
    }
    
    generate_report(kpis, region_summary, product_summary, monthly_summary, df, config)
    
    assert test_output.exists(), "Excel dosyası oluşturulamadı."
    wb = load_workbook(test_output)
    assert "Dashboard" in wb.sheetnames, "Dashboard sayfası eksik."
    assert "Detail Data" in wb.sheetnames, "Detail Data sayfası eksik."