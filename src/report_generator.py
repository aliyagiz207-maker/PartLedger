import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

def generate_report(
    kpis,
    region_summary,
    product_summary,
    monthly_summary,
    df,
    config,
):
    KPI_LABELS_TR = {
        "Total Quantity": "Toplam Adet",
        "Total Revenue": "Toplam Ciro",
        "Total Cost": "Toplam Maliyet",
        "Total Profit": "Toplam Kâr",
        "Profit Margin": "Kâr Marjı",
    }

    if getattr(sys, "frozen", False):
        application_root = Path(sys.executable).resolve().parent
        resource_root = Path(sys._MEIPASS)
    else:
        application_root = Path(__file__).resolve().parent.parent
        resource_root = application_root

    output_file = application_root / config["output_file"]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dashboard"

    header_fill = PatternFill(fill_type="solid", start_color="1F4E78", end_color="1F4E78")
    section_fill = PatternFill(fill_type="solid", start_color="D9EAF7", end_color="D9EAF7")
    kpi_fill = PatternFill(fill_type="solid", start_color="EAF2F8", end_color="EAF2F8")
    white_fill = PatternFill(fill_type="solid", start_color="FFFFFF", end_color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.freeze_panes = "A5"

    widths = {"A": 20, "B": 17, "C": 17, "D": 17, "E": 17, "F": 17, "G": 17, "H": 17, "I": 17, "J": 17, "K": 4, "L": 17, "M": 17, "N": 17}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    logo_file = resource_root / config["logo_path"]
    if logo_file.exists():
        logo = Image(logo_file)
        logo.width = 150
        logo.height = 75
        sheet.add_image(logo, "L1")

    sheet.merge_cells("A1:J1")
    sheet["A1"] = config["dashboard_title"]
    sheet["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    sheet["A1"].fill = header_fill
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 32

    sheet.merge_cells("A2:J2")
    sheet["A2"] = config["company_name"]
    sheet["A2"].font = Font(size=12, bold=True, color="404040")
    sheet["A2"].alignment = Alignment(vertical="center")

    if not df.empty and "Date" in df.columns:
        min_date = df["Date"].min()
        max_date = df["Date"].max()
        period_text = f"Analiz Dönemi: {min_date.strftime('%d.%m.%Y')} - {max_date.strftime('%d.%m.%Y')}"
        sheet.merge_cells("A3:J3")
        sheet["A3"] = period_text
        sheet["A3"].font = Font(size=10, italic=True, color="666666")

    kpi_positions = [("A5:B7", "A5", "A6"), ("C5:D7", "C5", "C6"), ("E5:F7", "E5", "E6"), ("G5:H7", "G5", "G6"), ("I5:J7", "I5", "I6")]
    kpi_items = list(kpis.items())

    for index, (key, value) in enumerate(kpi_items[:5]):
        merge_range, label_cell, value_cell = kpi_positions[index]
        sheet.merge_cells(merge_range)
        sheet.unmerge_cells(merge_range)
        start_col = sheet[label_cell].column
        start_letter = get_column_letter(start_col)
        end_letter = get_column_letter(start_col + 1)
        sheet.merge_cells(f"{start_letter}5:{end_letter}5")
        sheet.merge_cells(f"{start_letter}6:{end_letter}7")

        label = sheet[f"{start_letter}5"]
        value_cell_obj = sheet[f"{start_letter}6"]

        label.value = KPI_LABELS_TR.get(key, key)
        label.fill = kpi_fill
        label.font = Font(size=10, bold=True, color="44546A")
        label.alignment = Alignment(horizontal="center", vertical="center")
        label.border = thin_border

        value_cell_obj.fill = kpi_fill
        value_cell_obj.font = Font(size=16, bold=True, color="1F4E78")
        value_cell_obj.alignment = Alignment(horizontal="center", vertical="center")
        value_cell_obj.border = thin_border

        if isinstance(value, float):
            if "Margin" in key:
                value_cell_obj.value = value / 100
                value_cell_obj.number_format = "0.00%"
            else:
                value_cell_obj.value = value
                value_cell_obj.number_format = f'{config["currency"]}#,##0.00'
        else:
            value_cell_obj.value = value
            if "Quantity" not in key:
                value_cell_obj.number_format = f'{config["currency"]}#,##0'

    # --- BÖLGESEL PERFORMANS ---
    region_analysis = region_summary.sort_values("Revenue", ascending=False).reset_index(drop=True)

    sheet["A10"] = "Bölgesel Performans"
    sheet["A10"].font = Font(size=14, bold=True, color="1F1F1F")

    headers = ["Bölge", "Ciro", "Kâr", "Marj"]
    for column_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=11, column=column_index)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    region_row = 12
    for _, row_data in region_analysis.iterrows():
        sheet.cell(row=region_row, column=1, value=row_data["Region"])
        sheet.cell(row=region_row, column=2, value=row_data["Revenue"])
        sheet.cell(row=region_row, column=3, value=row_data["Profit"])
        sheet.cell(row=region_row, column=4, value=row_data["Margin"] / 100)

        for column_index in range(1, 5):
            cell = sheet.cell(row=region_row, column=column_index)
            cell.border = thin_border
            if column_index > 1:
                cell.alignment = Alignment(horizontal="right")

        sheet.cell(row=region_row, column=2).number_format = f'{config["currency"]}#,##0'
        sheet.cell(row=region_row, column=3).number_format = f'{config["currency"]}#,##0'
        sheet.cell(row=region_row, column=4).number_format = "0.00%"
        region_row += 1

    region_chart = BarChart()
    region_data_ref = Reference(sheet, min_col=2, min_row=11, max_row=region_row - 1)
    region_categories = Reference(sheet, min_col=1, min_row=12, max_row=region_row - 1)
    region_chart.add_data(region_data_ref, titles_from_data=True)
    region_chart.set_categories(region_categories)
    region_chart.type = "bar"
    region_chart.style = 10
    region_chart.title = "Bölgeye Göre Ciro"
    region_chart.y_axis.title = "Bölge"
    region_chart.x_axis.title = "Ciro"
    region_chart.width = 13
    region_chart.height = 6.5
    region_chart.legend = None
    
    # ETIKET DUZELTMELERI
    region_chart.dataLabels = DataLabelList()
    region_chart.dataLabels.showVal = True
    region_chart.dataLabels.showSerName = False
    region_chart.dataLabels.showCatName = False
    region_chart.dataLabels.showLegendKey = False
    
    region_chart.gapWidth = 60
    sheet.add_chart(region_chart, "F10")

    # --- EN ÇOK SATAN 5 ÜRÜN ---
    product_analysis = product_summary.sort_values("Revenue", ascending=False).head(5).reset_index(drop=True)

    sheet["A20"] = "En Çok Satan 5 Ürün"
    sheet["A20"].font = Font(size=14, bold=True, color="1F1F1F")

    for column_index, header in enumerate(["Ürün", "Ciro", "Kâr", "Marj"], start=1):
        cell = sheet.cell(row=21, column=column_index)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    product_row = 22
    for _, row_data in product_analysis.iterrows():
        sheet.cell(row=product_row, column=1, value=row_data["Product"])
        sheet.cell(row=product_row, column=2, value=row_data["Revenue"])
        sheet.cell(row=product_row, column=3, value=row_data["Profit"])
        sheet.cell(row=product_row, column=4, value=row_data["Margin"] / 100)

        for column_index in range(1, 5):
            cell = sheet.cell(row=product_row, column=column_index)
            cell.border = thin_border
            if column_index > 1:
                cell.alignment = Alignment(horizontal="right")

        sheet.cell(row=product_row, column=2).number_format = f'{config["currency"]}#,##0'
        sheet.cell(row=product_row, column=3).number_format = f'{config["currency"]}#,##0'
        sheet.cell(row=product_row, column=4).number_format = "0.00%"
        product_row += 1

    product_chart = BarChart()
    product_data_ref = Reference(sheet, min_col=2, min_row=21, max_row=product_row - 1)
    product_categories = Reference(sheet, min_col=1, min_row=22, max_row=product_row - 1)
    product_chart.add_data(product_data_ref, titles_from_data=True)
    product_chart.set_categories(product_categories)
    product_chart.type = "bar"
    product_chart.style = 10
    product_chart.title = "Ciroya Göre En Çok Satan 5 Ürün"
    product_chart.y_axis.title = "Ürün"
    product_chart.x_axis.title = "Ciro"
    product_chart.width = 13
    product_chart.height = 6.5
    product_chart.legend = None
    
    # ETIKET DUZELTMELERI
    product_chart.dataLabels = DataLabelList()
    product_chart.dataLabels.showVal = True
    product_chart.dataLabels.showSerName = False
    product_chart.dataLabels.showCatName = False
    product_chart.dataLabels.showLegendKey = False
    
    product_chart.gapWidth = 60
    sheet.add_chart(product_chart, "F20")

    # --- AYLIK PERFORMANS ---
    monthly_analysis = monthly_summary.sort_values("Month").reset_index(drop=True)

    sheet["A30"] = "Aylık Performans"
    sheet["A30"].font = Font(size=14, bold=True, color="1F1F1F")

    for column_index, header in enumerate(["Ay", "Ciro", "Kâr"], start=1):
        cell = sheet.cell(row=31, column=column_index)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    monthly_row = 32
    for _, row_data in monthly_analysis.iterrows():
        sheet.cell(row=monthly_row, column=1, value=row_data["Month"])
        sheet.cell(row=monthly_row, column=2, value=row_data["Revenue"])
        sheet.cell(row=monthly_row, column=3, value=row_data["Profit"])

        for column_index in range(1, 4):
            cell = sheet.cell(row=monthly_row, column=column_index)
            cell.border = thin_border
            if column_index > 1:
                cell.alignment = Alignment(horizontal="right")

        sheet.cell(row=monthly_row, column=2).number_format = f'{config["currency"]}#,##0'
        sheet.cell(row=monthly_row, column=3).number_format = f'{config["currency"]}#,##0'
        monthly_row += 1

    monthly_chart = BarChart()
    monthly_data_ref = Reference(sheet, min_col=2, min_row=31, max_row=monthly_row - 1)
    monthly_categories = Reference(sheet, min_col=1, min_row=32, max_row=monthly_row - 1)
    monthly_chart.add_data(monthly_data_ref, titles_from_data=True)
    monthly_chart.set_categories(monthly_categories)
    monthly_chart.type = "col"
    monthly_chart.style = 10
    monthly_chart.title = "Aylık Ciro"
    monthly_chart.y_axis.title = "Ciro"
    monthly_chart.x_axis.title = "Ay"
    monthly_chart.width = 13
    monthly_chart.height = 6.5
    monthly_chart.legend = None
    
    # ETIKET DUZELTMELERI
    monthly_chart.dataLabels = DataLabelList()
    monthly_chart.dataLabels.showVal = True
    monthly_chart.dataLabels.showSerName = False
    monthly_chart.dataLabels.showCatName = False
    monthly_chart.dataLabels.showLegendKey = False
    
    monthly_chart.gapWidth = 60
    monthly_chart.y_axis.scaling.min = 0
    sheet.add_chart(monthly_chart, "F30")

    # --- DETAIL DATA (HAM VERİ) SEKMESİ ---
    detail_sheet = workbook.create_sheet("Detail Data")
    detail_sheet.sheet_view.showGridLines = False
    detail_sheet.freeze_panes = "A2"

    for row_data in dataframe_to_rows(df, index=False, header=True):
        detail_sheet.append(row_data)

    detail_sheet.auto_filter.ref = detail_sheet.dimensions

    date_column_index = None
    for column_index, header_cell in enumerate(detail_sheet[1], start=1):
        if header_cell.value == "Date":
            date_column_index = column_index
            break

    if date_column_index is not None:
        date_letter = get_column_letter(date_column_index)
        for row_index in range(2, detail_sheet.max_row + 1):
            detail_sheet[f"{date_letter}{row_index}"].number_format = "dd.mm.yyyy"
            
    currency_columns = ["UnitPrice", "UnitCost", "Revenue", "Profit"]
    for col_idx, header_cell in enumerate(detail_sheet[1], start=1):
        if header_cell.value in currency_columns:
            col_letter = get_column_letter(col_idx)
            for row_idx in range(2, detail_sheet.max_row + 1):
                detail_sheet[f"{col_letter}{row_idx}"].number_format = f'{config["currency"]}#,##0.00'

    for cell in detail_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for column_cells in detail_sheet.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        detail_sheet.column_dimensions[column_cells[0].column_letter].width = min(length + 3, 30)

    detail_sheet.sheet_view.zoomScale = 90

    # --- SAYFA VE YAZDIRMA AYARLARI ---
    sheet.print_area = "A1:N53"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5

    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)